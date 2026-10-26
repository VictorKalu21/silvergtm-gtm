#!/usr/bin/env python3
"""build-xlsx.py :: one workbook from assemble.js's deliverable/ CSVs (UK home-extension builders; modelled on the
pools run's build-xlsx.py). Tabs: Summary, Qualified, Review, Excluded, All, By Nation.
Usage: python3 build-xlsx.py [--deliverable <dir>]  (default: <run>/deliverable)"""
import csv, json, os, argparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser(); ap.add_argument("--deliverable", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "deliverable")); a = ap.parse_args()
D = a.deliverable
HEAD = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="1F4E78")
NUM = {"google_rating", "google_reviews", "location_count", "latitude", "longitude"}

def add_table(wb, title, path, max_rows=None):
    ws = wb.create_sheet(title)
    if not os.path.exists(path): ws.append(["(missing)", path]); return ws
    with open(path, newline="", encoding="utf-8") as f:
        rd = csv.reader(f); cols = next(rd, [])
        ws.append(cols)
        for c in ws[1]: c.font = HEAD; c.fill = FILL; c.alignment = Alignment(vertical="center")
        n = 0
        for row in rd:
            out = []
            for h, v in zip(cols, row):
                if h in NUM and v not in ("", None):
                    try: v = float(v) if "." in v else int(v)
                    except ValueError: pass
                out.append(v)
            ws.append(out); n += 1
            if max_rows and n >= max_rows: break
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(cols, 1):
        w = 14
        if h in ("company_name", "full_address", "website", "google_types", "services_offered", "verdict_reason", "work_types"): w = 36
        if h in ("build_evidence", "evidence_source", "google_maps_url", "facebook_url"): w = 44
        if h in ("verdict", "priority", "nation", "postcode", "phone", "email_source", "residential", "qualify_track"): w = 11
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws

wb = Workbook(); ws = wb.active; ws.title = "Summary"
s = json.load(open(os.path.join(D, "summary.json"))) if os.path.exists(os.path.join(D, "summary.json")) else {}
rows = [("UK home-extension builders — prospect list summary", ""), ("generated", s.get("generated", "")), ("", ""),
        ("Total unique businesses collected", s.get("total_unique_collected")), ("Excluded at the Maps stage", s.get("maps_stage_excluded")),
        ("Reached the site-text gate", s.get("unique_businesses_after_maps_qualify")), ("QUALIFIED extension builders", s.get("qualified")),
        ("Review", s.get("review")), ("Excluded at the site-text gate", s.get("excluded_site_text")), ("In-run duplicates collapsed", s.get("duplicates_collapsed_phone_or_name_plus_street")), ("", "")]
rows += [("Review rows by reason", "")] + [(k, v) for k, v in (s.get("review_by_reason") or {}).items()] + [("", "")]
rows += [("Work types named on Qualified sites", "")] + [(k, v) for k, v in (s.get("work_types_qualified") or {}).items()] + [("", "")]
rows += [("Qualified priority (A = 25+ reviews, 4.3+, website; B = 5+ reviews; C = rest)", "")] + [(k, v) for k, v in (s.get("qualified_priority") or {}).items()] + [("", "")]
rows += [("Field fill, Qualified rows (%)", "")] + [(k, v) for k, v in (s.get("fill_qualified_pct") or {}).items()] + [("", "")]
rows += [("Website fetch status", "")] + [(k, v) for k, v in (s.get("website_fetch_status") or {}).items()]
for r in rows: ws.append(list(r))
ws["A1"].font = Font(bold=True, size=14); ws.column_dimensions["A"].width = 60; ws.column_dimensions["B"].width = 18
for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
    if r[0].value and r[1].value == "" and r[0].value != "": r[0].font = Font(bold=True)
add_table(wb, "Qualified", os.path.join(D, "extension_builders_qualified.csv"))
add_table(wb, "Review", os.path.join(D, "extension_builders_review.csv"))
add_table(wb, "Excluded", os.path.join(D, "extension_builders_excluded.csv"))
add_table(wb, "All", os.path.join(D, "extension_builders_all.csv"))
bs = wb.create_sheet("By Nation"); bs.append(["nation", "reached site gate", "qualified", "review"])
for c in bs[1]: c.font = HEAD; c.fill = FILL
for st, v in sorted((s.get("by_nation") or {}).items(), key=lambda kv: -kv[1].get("qualified", 0)): bs.append([st, v.get("site_text_rows"), v.get("qualified"), v.get("review")])
bs.freeze_panes = "A2"
out = os.path.join(D, "extension_builders.xlsx"); wb.save(out); print("wrote", out, "| sheets:", wb.sheetnames)
